#!/usr/bin/env python3
"""
Import company-level YUCG prospect targets from the spreadsheet into SQLite.

Source of truth: data/YUCG_Prospect_List.xlsx (see data/README.md).

Usage (from repo root):
  python scripts/sync_prospects_from_xlsx.py
  python scripts/sync_prospects_from_xlsx.py --xlsx data/YUCG_Prospect_List.expanded.xlsx
  python scripts/sync_prospects_from_xlsx.py --dry-run
"""

from __future__ import annotations

import argparse
import json
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = REPO_ROOT / "data" / "YUCG_Prospect_List.xlsx"
DB_PATH = REPO_ROOT / "backend" / "clientreach.db"
SCHEMA_PATH = REPO_ROOT / "data" / "prospects.schema.json"

# Spreadsheet header -> yucg_prospect_targets column
KNOWN_COLUMN_MAP: dict[str, str] = {
    "Company": "company",
    "Sector": "sector",
    "Why Attractive Prospect for YUCG": "why_attractive",
    "Suggested Engagement Theme": "suggested_engagement_theme",
    "Yale / YUCG Hook": "yale_yucg_hook",
}

# Spreadsheet headers not yet promoted to dedicated DB columns (Agent 2+ extensions)
EXTRA_JSON_KEYS: dict[str, str] = {
    "Outreach Priority": "outreach_priority",
    "Contact Type": "contact_type",
    "Target Role Title": "target_role_title",
    "Incentive Score": "incentive_score",
    "Score Rationale": "score_rationale",
    "Verification Source URL": "verification_source_url",
    "Recommended First Message Angle": "first_message_angle",
    "Contact Discovery Hint": "discovery_hint",
}

TABLE_DDL = """
CREATE TABLE IF NOT EXISTS yucg_prospect_targets (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    company TEXT NOT NULL UNIQUE,
    sector TEXT,
    why_attractive TEXT,
    suggested_engagement_theme TEXT,
    yale_yucg_hook TEXT,
    source_row INTEGER,
    extra_json TEXT,
    synced_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
);
CREATE INDEX IF NOT EXISTS idx_yucg_prospect_targets_sector ON yucg_prospect_targets(sector);
"""


def _cell_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _load_headers_and_rows(xlsx_path: Path) -> tuple[list[str], list[tuple[int, dict[str, object]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("Missing openpyxl. Install with: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in header_row]

    rows: list[tuple[int, dict[str, object]]] = []
    for excel_row_idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in values):
            continue
        record = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers))}
        company = _cell_str(record.get("Company"))
        if not company:
            continue
        rows.append((excel_row_idx, record))
    wb.close()
    return headers, rows


def _row_to_db_payload(source_row: int, record: dict[str, object]) -> tuple[dict[str, object], dict[str, object]]:
    mapped: dict[str, object] = {"source_row": source_row}
    extra: dict[str, object] = {}

    for header, value in record.items():
        if not header:
            continue
        if header in KNOWN_COLUMN_MAP:
            field = KNOWN_COLUMN_MAP[header]
            mapped[field] = _cell_str(value)
        elif header in EXTRA_JSON_KEYS:
            key = EXTRA_JSON_KEYS[header]
            extra[key] = value if isinstance(value, (int, float, bool)) else _cell_str(value)
        else:
            # Unknown column (future Agent 2+ additions): preserve under raw header
            extra[header] = value if isinstance(value, (int, float, bool)) else _cell_str(value)

    if not mapped.get("company"):
        raise ValueError(f"Row {source_row} missing Company")

    return mapped, extra


def ensure_table(conn: sqlite3.Connection) -> None:
    conn.executescript(TABLE_DDL)


def sync_prospects(
    xlsx_path: Path,
    db_path: Path,
    *,
    dry_run: bool = False,
) -> dict[str, int | list[str]]:
    headers, rows = _load_headers_and_rows(xlsx_path)
    unknown_headers = [
        h
        for h in headers
        if h and h not in KNOWN_COLUMN_MAP and h not in EXTRA_JSON_KEYS
    ]

    stats: dict[str, int | list[str]] = {
        "inserted": 0,
        "updated": 0,
        "skipped": 0,
        "total_rows": len(rows),
        "headers": headers,
        "unknown_headers": unknown_headers,
    }

    if dry_run:
        return stats

    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    try:
        ensure_table(conn)
        now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

        for source_row, record in rows:
            try:
                mapped, extra = _row_to_db_payload(source_row, record)
            except ValueError:
                stats["skipped"] = int(stats["skipped"]) + 1
                continue

            company = mapped["company"]
            extra_json = json.dumps(extra, ensure_ascii=False) if extra else None

            existing = conn.execute(
                "SELECT id FROM yucg_prospect_targets WHERE company = ?",
                (company,),
            ).fetchone()

            if existing:
                conn.execute(
                    """UPDATE yucg_prospect_targets SET
                        sector = ?, why_attractive = ?, suggested_engagement_theme = ?,
                        yale_yucg_hook = ?, source_row = ?, extra_json = ?,
                        synced_at = ?, updated_at = CURRENT_TIMESTAMP
                    WHERE company = ?""",
                    (
                        mapped.get("sector"),
                        mapped.get("why_attractive"),
                        mapped.get("suggested_engagement_theme"),
                        mapped.get("yale_yucg_hook"),
                        mapped.get("source_row"),
                        extra_json,
                        now,
                        company,
                    ),
                )
                stats["updated"] = int(stats["updated"]) + 1
            else:
                conn.execute(
                    """INSERT INTO yucg_prospect_targets (
                        company, sector, why_attractive, suggested_engagement_theme,
                        yale_yucg_hook, source_row, extra_json, synced_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        company,
                        mapped.get("sector"),
                        mapped.get("why_attractive"),
                        mapped.get("suggested_engagement_theme"),
                        mapped.get("yale_yucg_hook"),
                        mapped.get("source_row"),
                        extra_json,
                        now,
                    ),
                )
                stats["inserted"] = int(stats["inserted"]) + 1

        conn.commit()
    finally:
        conn.close()

    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description="Sync YUCG prospect spreadsheet into SQLite")
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=DEFAULT_XLSX,
        help=f"Path to prospect list xlsx (default: {DEFAULT_XLSX.relative_to(REPO_ROOT)})",
    )
    parser.add_argument(
        "--db",
        type=Path,
        default=DB_PATH,
        help=f"SQLite database path (default: {DB_PATH.relative_to(REPO_ROOT)})",
    )
    parser.add_argument("--dry-run", action="store_true", help="Parse xlsx only; do not write to DB")
    args = parser.parse_args()

    if not args.xlsx.is_file():
        print(f"Spreadsheet not found: {args.xlsx}", file=sys.stderr)
        sys.exit(1)

    if not args.dry_run and not args.db.parent.is_dir():
        print(f"Database directory missing: {args.db.parent}", file=sys.stderr)
        sys.exit(1)

    stats = sync_prospects(args.xlsx, args.db, dry_run=args.dry_run)

    print(f"Source: {args.xlsx}")
    print(f"Headers ({len(stats['headers'])}): {', '.join(stats['headers'])}")
    if stats["unknown_headers"]:
        print(f"Unknown headers (stored in extra_json): {', '.join(stats['unknown_headers'])}")
    print(f"Rows parsed: {stats['total_rows']}")
    if args.dry_run:
        print("Dry run — no database changes.")
    else:
        print(f"Database: {args.db}")
        print(f"Inserted: {stats['inserted']}, Updated: {stats['updated']}, Skipped: {stats['skipped']}")
        if SCHEMA_PATH.is_file():
            print(f"Schema: {SCHEMA_PATH.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
