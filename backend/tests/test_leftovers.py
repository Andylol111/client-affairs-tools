"""Bounce + Think-Cell pack checks. From backend/: python tests/test_leftovers.py"""
from __future__ import annotations

import os
import sys
from io import BytesIO
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
os.environ.setdefault("JWT_SECRET", "p2-test-secret-not-a-known-default-xx")

from app.jwt_utils import COOKIE_NAME, create_token
from app.services.gmail_reply_sync import thread_has_bounce
from app.services.thinkcell_pack import build_pack_bytes


def test_bounce_from_mailer_daemon() -> None:
    thread = {
        "messages": [
            {
                "payload": {
                    "headers": [
                        {"name": "From", "value": "Mail Delivery Subsystem <mailer-daemon@googlemail.com>"},
                        {"name": "Subject", "value": "Undeliverable: Quick question"},
                    ]
                }
            }
        ]
    }
    assert thread_has_bounce(thread) is True
    assert thread_has_bounce({"messages": []}) is False


def test_thinkcell_table_names() -> None:
    from openpyxl import load_workbook

    data = build_pack_bytes([["Acme", "Aviation", "ASD", 8, "pending"]])
    wb = load_workbook(BytesIO(data))
    names = {t for ws in wb.worksheets for t in ws.tables}
    assert names == {"Table_Slate", "Table_Addresses", "Table_Pipeline", "Table_Send"}


def test_pending_2fa_not_authenticated() -> None:
    from fastapi.testclient import TestClient
    from main import app

    token = create_token(1, "test@yale.edu", extra={"2fa": "pending"}, expiry_hours=0.25)
    with TestClient(app) as client:
        r = client.get("/api/auth/me", cookies={COOKIE_NAME: token})
        assert r.status_code == 200
        assert r.json().get("authenticated") is False


def test_anthropic_catalog_opus_to_haiku() -> None:
    from app.services.llm import BEDROCK_ANTHROPIC, is_bedrock_model, list_models

    ids = {m["id"] for m in BEDROCK_ANTHROPIC}
    labels = {m["label"] for m in BEDROCK_ANTHROPIC}
    assert ids == {
        "us.anthropic.claude-opus-5",
        "us.anthropic.claude-sonnet-5",
        "us.anthropic.claude-haiku-4-5-20251001-v1:0",
    }
    assert "Claude Opus 4" not in labels
    assert is_bedrock_model("us.anthropic.claude-opus-5")
    assert not is_bedrock_model("ollama:llama3.2")
    catalog = list_models()
    assert [g["id"] for g in catalog["groups"]] == ["anthropic"]
    assert not any(m["id"].startswith("ollama:") for g in catalog["groups"] for m in g["models"])


def test_spa_week_route_uses_index() -> None:
    import tempfile

    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from fastapi.responses import FileResponse
    from main import spa_file

    with tempfile.TemporaryDirectory() as d:
        root = Path(d)
        (root / "index.html").write_text("spa", encoding="utf-8")
        (root / "favicon.ico").write_text("ico", encoding="utf-8")
        assert spa_file(root, "yucgoutreach").name == "index.html"
        assert spa_file(root, "login").name == "index.html"
        assert spa_file(root, "favicon.ico").name == "favicon.ico"

        mini = FastAPI()

        @mini.get("/{full_path:path}")
        async def spa_fallback(full_path: str):
            return FileResponse(spa_file(root, full_path))

        with TestClient(mini) as client:
            r = client.get("/yucgoutreach")
            assert r.status_code == 200
            assert r.text == "spa"


def test_health_ok_without_spa() -> None:
    from fastapi.testclient import TestClient
    from main import app

    with TestClient(app) as client:
        r = client.get("/api/health")
        assert r.status_code == 200
        assert r.json()["status"] == "ok"


if __name__ == "__main__":
    test_bounce_from_mailer_daemon()
    test_thinkcell_table_names()
    test_pending_2fa_not_authenticated()
    test_anthropic_catalog_opus_to_haiku()
    test_spa_week_route_uses_index()
    print("ok")
