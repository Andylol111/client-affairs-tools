"""YUCG_Outreach_Pack.xlsx — stable Think-Cell table names. Do not rename tables."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

TABLE_SLATE = ("company", "sector", "contact_type", "incentive", "find_status")
TABLE_ADDRESSES = ("name", "title", "email", "verify_status", "source")
TABLE_PIPELINE = ("cold", "contacted", "replied", "meeting", "closed")
TABLE_SEND = ("queued", "sent_today", "errors", "open_rate", "reply_rate")

LOCAL_PACK = Path(__file__).resolve().parents[3] / "data" / "YUCG_Outreach_Pack.xlsx"
LOCAL_SLATE_DIR = Path(__file__).resolve().parents[3] / "data" / "releases"


def _write_table(ws, name: str, headers: tuple[str, ...], rows: list[list]) -> None:
    ws.append(list(headers))
    body = rows or [[""] * len(headers)]
    for row in body:
        padded = list(row) + [""] * (len(headers) - len(row))
        ws.append(padded[: len(headers)])
    last_row = 1 + max(1, len(body))
    last_col = get_column_letter(len(headers))
    table = Table(displayName=name, ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def build_pack_bytes(
    slate_rows: list[list] | None = None,
    address_rows: list[list] | None = None,
    pipeline_row: list | None = None,
    send_row: list | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Slate"
    _write_table(ws, "Table_Slate", TABLE_SLATE, slate_rows or [])
    ws2 = wb.create_sheet("Addresses")
    _write_table(ws2, "Table_Addresses", TABLE_ADDRESSES, address_rows or [])
    ws3 = wb.create_sheet("Pipeline")
    _write_table(ws3, "Table_Pipeline", TABLE_PIPELINE, [pipeline_row or [0, 0, 0, 0, 0]])
    ws4 = wb.create_sheet("Send")
    _write_table(ws4, "Table_Send", TABLE_SEND, [send_row or [0, 0, 0, 0, 0]])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def pack_rows_for_release(db, release_id: int) -> dict:
    cur = await db.execute(
        "SELECT company, sector, contact_type, incentive_score, find_status FROM outreach_release_targets WHERE release_id = ? ORDER BY id",
        (release_id,),
    )
    slate = [
        [r["company"], r["sector"], r["contact_type"], r["incentive_score"], r["find_status"]]
        for r in await cur.fetchall()
    ]
    cur = await db.execute(
        "SELECT full_name, title, email, email_status, source_url FROM outreach_release_people WHERE release_id = ? ORDER BY id",
        (release_id,),
    )
    addresses = [
        [r["full_name"], r["title"], r["email"], r["email_status"], r["source_url"]]
        for r in await cur.fetchall()
    ]
    cur = await db.execute(
        """SELECT
             SUM(CASE WHEN COALESCE(c.pipeline_status, 'cold') = 'cold' THEN 1 ELSE 0 END) AS cold,
             SUM(CASE WHEN c.pipeline_status = 'contacted' THEN 1 ELSE 0 END) AS contacted,
             SUM(CASE WHEN c.pipeline_status = 'replied' THEN 1 ELSE 0 END) AS replied,
             SUM(CASE WHEN c.pipeline_status = 'meeting' THEN 1 ELSE 0 END) AS meeting,
             SUM(CASE WHEN c.pipeline_status = 'closed' THEN 1 ELSE 0 END) AS closed
           FROM outreach_release_people p
           LEFT JOIN contacts c ON c.id = p.contact_id
           WHERE p.release_id = ? AND p.kept = 1""",
        (release_id,),
    )
    pipe = await cur.fetchone()
    pipeline = [
        int(pipe["cold"] or 0) if pipe else 0,
        int(pipe["contacted"] or 0) if pipe else 0,
        int(pipe["replied"] or 0) if pipe else 0,
        int(pipe["meeting"] or 0) if pipe else 0,
        int(pipe["closed"] or 0) if pipe else 0,
    ]
    cur = await db.execute(
        """SELECT
             SUM(CASE WHEN cc.status = 'pending' THEN 1 ELSE 0 END) AS queued,
             SUM(CASE WHEN cc.status = 'sent' THEN 1 ELSE 0 END) AS sent_today,
             SUM(CASE WHEN cc.status IN ('error', 'bounced') THEN 1 ELSE 0 END) AS errors,
             SUM(CASE WHEN cc.opened_at IS NOT NULL THEN 1 ELSE 0 END) AS opened,
             SUM(CASE WHEN cc.replied_at IS NOT NULL THEN 1 ELSE 0 END) AS replied,
             COUNT(*) AS n
           FROM campaign_contacts cc
           WHERE cc.contact_id IN (
             SELECT contact_id FROM outreach_release_people
             WHERE release_id = ? AND kept = 1 AND contact_id IS NOT NULL
           )""",
        (release_id,),
    )
    send = await cur.fetchone()
    n = int(send["n"] or 0) if send else 0
    opened = int(send["opened"] or 0) if send else 0
    replied = int(send["replied"] or 0) if send else 0
    send_row = [
        int(send["queued"] or 0) if send else 0,
        int(send["sent_today"] or 0) if send else 0,
        int(send["errors"] or 0) if send else 0,
        round(opened / n, 3) if n else 0,
        round(replied / n, 3) if n else 0,
    ]
    return {"slate": slate, "addresses": addresses, "pipeline": pipeline, "send": send_row}


def write_week_slate(release_id: int, name: str, rows: list[dict]) -> dict:
    """Sibling workbook so people who never open the app still see this week's companies."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Week slate"
    ws.append(["company", "sector", "contact_type", "incentive", "source_url"])
    for row in rows:
        ws.append(
            [
                row.get("company"),
                row.get("sector"),
                row.get("contact_type"),
                row.get("incentive_score"),
                row.get("verification_source_url"),
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    LOCAL_SLATE_DIR.mkdir(parents=True, exist_ok=True)
    path = LOCAL_SLATE_DIR / f"week_slate_{release_id}.xlsx"
    path.write_bytes(data)
    out = {"path": str(path), "bytes": len(data)}
    from app.services.object_catalog import catalog_bucket, put_bytes, record_sync
    from app.services.graph_onedrive import configured, upload_bytes

    key = f"prospects/releases/{release_id}.xlsx"
    if catalog_bucket():
        meta = put_bytes(
            key,
            data,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        record_sync("week_slate", meta["key"], meta["sha256"], meta["size"])
        out["s3"] = meta["key"]
    if configured():
        try:
            out["graph"] = upload_bytes(f"YUCG_week_slate_{release_id}.xlsx", data)
        except Exception:
            pass
    return out


def write_pack_bytes(data: bytes) -> dict:
    LOCAL_PACK.parent.mkdir(parents=True, exist_ok=True)
    LOCAL_PACK.write_bytes(data)
    out = {"path": str(LOCAL_PACK), "bytes": len(data)}
    from app.services.object_catalog import catalog_bucket, put_bytes, record_sync

    if catalog_bucket():
        meta = put_bytes(
            "packs/YUCG_Outreach_Pack.xlsx",
            data,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        record_sync(
            "pack",
            meta["key"],
            meta["sha256"],
            meta["size"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        out["s3"] = meta["key"]
    return out


async def rebuild_release_pack(db, release_id: int) -> dict:
    rows = await pack_rows_for_release(db, release_id)
    data = build_pack_bytes(rows["slate"], rows["addresses"], rows["pipeline"], rows["send"])
    written = write_pack_bytes(data)
    from app.services.graph_onedrive import upload_pack_if_configured

    graph = upload_pack_if_configured(data)
    if graph:
        written["graph"] = graph
    return written
