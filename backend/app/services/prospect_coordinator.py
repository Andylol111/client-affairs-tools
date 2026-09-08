"""
Outreach coordinator — loads YUCG prospect targets from data/YUCG_Prospect_List.xlsx
with file-mtime caching. Scoring and filters align with expanded spreadsheet columns
when present (Agent 2); falls back to heuristics on base columns.
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Repo root: backend/app/services -> backend -> repo
_REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_XLSX_PATH = _REPO_ROOT / "data" / "YUCG_Prospect_List.xlsx"
SHEET_NAME = "YUCG Prospects"

# Normalized header -> canonical field
_FIELD_ALIASES: dict[str, str] = {
    "company": "company",
    "sector": "sector",
    "why_attractive_prospect_for_yucg": "why_attractive",
    "why_attractive": "why_attractive",
    "suggested_engagement_theme": "engagement_theme",
    "engagement_theme": "engagement_theme",
    "yale_yucg_hook": "yale_hook",
    "yale_hook": "yale_hook",
    "outreach_priority": "outreach_priority",
    "priority": "outreach_priority",
    "contact_type": "contact_type",
    "target_role_title": "target_role_title",
    "incentive_score": "incentive_score",
    "verification_source_url": "verification_source_url",
    "recommended_first_message_angle": "first_message_angle",
    "contact_discovery_hint": "discovery_hint",
    "score_rationale": "score_rationale",
}

YUCG_SERVICE_KEYWORDS: list[tuple[str, list[str]]] = [
    ("Marketing & Branding", ["marketing", "brand", "positioning", "narrative", "campaign", "go-to-market", "gtm"]),
    ("Market Entry & Expansion", ["entry", "expansion", "market entry", "geographic", "launch", "international", "us market"]),
    ("Financial Strategy", ["financial", "tco", "roi", "pricing", "revenue", "unit economics", "fleet planning", "capex"]),
    ("Operations Improvement", ["operations", "ops", "efficiency", "supply chain", "logistics", "process"]),
    ("Product Development", ["product", "roadmap", "feature", "platform", "innovation"]),
    ("Data Analysis & Research", ["data", "research", "analytics", "survey", "gen-z", "gen z", "student market", "modeling"]),
]

_DEFAULT_WEIGHTS = {
    "incentive": 0.45,
    "priority": 0.25,
    "yale_hook": 0.15,
    "contact_type_match": 0.15,
}

_cache: dict[str, Any] = {"mtime": None, "rows": [], "path": None}


def prospect_xlsx_path() -> Path:
    import os

    env = (os.getenv("YUCG_PROSPECT_XLSX") or "").strip()
    if env:
        p = Path(env)
        return p if p.is_absolute() else _REPO_ROOT / env
    if (os.getenv("CATALOG_BUCKET") or "").strip():
        return _xlsx_from_catalog()
    return DEFAULT_XLSX_PATH


def _xlsx_from_catalog() -> Path:
    cache = _REPO_ROOT / "data" / ".cache" / "current.xlsx"
    cache.parent.mkdir(parents=True, exist_ok=True)
    if cache.exists():
        return cache
    from app.services.object_catalog import get_bytes

    cache.write_bytes(get_bytes("prospects/current.xlsx"))
    return cache


def _normalize_header(raw: str | None) -> str:
    if not raw:
        return ""
    return re.sub(r"[^a-z0-9]+", "_", str(raw).strip().lower()).strip("_")


def _coerce_float(val: Any) -> float | None:
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _coerce_int(val: Any) -> int | None:
    f = _coerce_float(val)
    if f is None:
        return None
    return int(round(f))


def _text(val: Any) -> str:
    if val is None:
        return ""
    return str(val).strip()


def derive_yucg_service_tags(engagement_theme: str, why_attractive: str = "") -> list[str]:
    blob = f"{engagement_theme} {why_attractive}".lower()
    tags: list[str] = []
    for label, keywords in YUCG_SERVICE_KEYWORDS:
        if any(kw in blob for kw in keywords):
            tags.append(label)
    if not tags:
        tags.append("Strategy & Consulting")
    return tags


def _default_verification_url(company: str, explicit: str) -> str:
    if explicit:
        return explicit
    q = company.replace(" ", "+")
    return f"https://www.linkedin.com/search/results/people/?keywords={q}"


def _heuristic_incentive(row: dict[str, Any]) -> float:
    why = _text(row.get("why_attractive")).lower()
    hook = _text(row.get("yale_hook"))
    score = 50.0
    if hook and hook.lower() not in ("none", "n/a", "-"):
        score += 18
    if "gen-z" in why or "gen z" in why or "student" in why:
        score += 12
    if "repeat" in why or "returning client" in why or "past client" in why:
        score += 10
    if "verif" in why or "public" in why or "linkedin" in why:
        score += 8
    if len(why) > 200:
        score += 5
    return min(100.0, max(0.0, score))


def _heuristic_priority(row: dict[str, Any]) -> int:
    incentive = row.get("incentive_score")
    if incentive is not None:
        if incentive >= 85:
            return 1
        if incentive >= 70:
            return 2
        if incentive >= 55:
            return 3
        if incentive >= 40:
            return 4
        return 5
    hook = _text(row.get("yale_hook"))
    if hook and hook.lower() not in ("none", "n/a"):
        return 2
    return 3


def _heuristic_contact_type(sector: str) -> str:
    s = sector.lower()
    if "airport" in s or "aviation" in s or "airline" in s:
        return "Head of Airport ASD"
    if "entertainment" in s or "studio" in s or "media" in s or "exhibition" in s:
        return "Studio Strategy"
    if "financial" in s or "bank" in s:
        return "CFO / Corporate Development"
    return "VP Business Development"


def _normalize_row(raw: dict[str, Any], row_index: int) -> dict[str, Any]:
    company = _text(raw.get("company"))
    sector = _text(raw.get("sector"))
    why = _text(raw.get("why_attractive"))
    theme = _text(raw.get("engagement_theme"))
    hook = _text(raw.get("yale_hook"))
    if hook.lower() in ("none", "n/a"):
        hook = ""

    incentive = _coerce_float(raw.get("incentive_score"))
    if incentive is None:
        incentive = _heuristic_incentive({"why_attractive": why, "yale_hook": hook})

    priority = _coerce_int(raw.get("outreach_priority"))
    if priority is None or priority < 1 or priority > 5:
        priority = _heuristic_priority({"incentive_score": incentive, "yale_hook": hook})

    contact_type = _text(raw.get("contact_type")) or _heuristic_contact_type(sector)
    verification = _text(raw.get("verification_source_url"))

    rationale = _text(raw.get("score_rationale"))
    if not rationale:
        parts = []
        if hook:
            parts.append("Yale/YUCG hook present")
        if incentive >= 75:
            parts.append("strong fit signals in prospect narrative")
        elif incentive >= 55:
            parts.append("moderate outreach fit")
        if not parts:
            parts.append("baseline sector opportunity")
        rationale = "; ".join(parts)

    return {
        "row_index": row_index,
        "company": company,
        "sector": sector,
        "why_attractive": why,
        "engagement_theme": theme,
        "yale_hook": hook or None,
        "has_yale_hook": bool(hook),
        "outreach_priority": priority,
        "contact_type": contact_type,
        "target_role_title": _text(raw.get("target_role_title")) or None,
        "incentive_score": round(incentive, 1),
        "verification_source_url": _default_verification_url(company, verification),
        "first_message_angle": _text(raw.get("first_message_angle")) or None,
        "discovery_hint": _text(raw.get("discovery_hint")) or None,
        "score_rationale": rationale,
        "yucg_service_tags": derive_yucg_service_tags(theme, why),
    }


def _load_rows_openpyxl(path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return []

        col_map: dict[int, str] = {}
        for idx, cell in enumerate(header_row):
            norm = _normalize_header(str(cell) if cell is not None else "")
            canonical = _FIELD_ALIASES.get(norm)
            if canonical:
                col_map[idx] = canonical

        out: list[dict[str, Any]] = []
        excel_row = 1
        for values in rows_iter:
            excel_row += 1
            raw: dict[str, Any] = {}
            for col_idx, field in col_map.items():
                if col_idx < len(values):
                    raw[field] = values[col_idx]
            if not _text(raw.get("company")):
                continue
            out.append(_normalize_row(raw, excel_row))
        return out
    finally:
        wb.close()


def _load_rows_pandas(path: Path) -> list[dict[str, Any]]:
    import pandas as pd

    sheet = SHEET_NAME
    try:
        xl = pd.ExcelFile(path, engine="openpyxl")
        if SHEET_NAME not in xl.sheet_names:
            sheet = xl.sheet_names[0]
        df = pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
    except ValueError:
        df = pd.read_excel(path, sheet_name=0, engine="openpyxl")

    if df.empty:
        return []

    col_map: dict[str, str] = {}
    for col in df.columns:
        norm = _normalize_header(str(col))
        canonical = _FIELD_ALIASES.get(norm)
        if canonical:
            col_map[str(col)] = canonical

    out: list[dict[str, Any]] = []
    for idx, row in df.iterrows():
        raw: dict[str, Any] = {}
        for src_col, field in col_map.items():
            val = row.get(src_col)
            if pd.isna(val):
                val = None
            raw[field] = val
        if not _text(raw.get("company")):
            continue
        # Excel row: header=1, data starts at 2; pandas index 0 -> row 2
        excel_row = int(idx) + 2
        out.append(_normalize_row(raw, excel_row))
    return out


def _load_rows_from_xlsx(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        raise FileNotFoundError(f"Prospect spreadsheet not found: {path}")

    try:
        return _load_rows_pandas(path)
    except ImportError:
        return _load_rows_openpyxl(path)


def invalidate_prospect_cache() -> None:
    _cache["mtime"] = None
    _cache["rows"] = []


def load_prospects(*, force_reload: bool = False) -> list[dict[str, Any]]:
    path = prospect_xlsx_path()
    mtime = path.stat().st_mtime if path.is_file() else None
    if (
        not force_reload
        and _cache.get("mtime") == mtime
        and _cache.get("path") == str(path)
        and _cache.get("rows")
    ):
        return list(_cache["rows"])

    rows = _load_rows_from_xlsx(path)
    _cache["mtime"] = mtime
    _cache["path"] = str(path)
    _cache["rows"] = rows
    return list(rows)


def filter_prospects(
    rows: list[dict[str, Any]] | None = None,
    *,
    sector: str | None = None,
    priority: int | None = None,
    min_incentive_score: float | None = None,
    contact_type: str | None = None,
    q: str | None = None,
    limit: int | None = None,
) -> list[dict[str, Any]]:
    result = list(rows) if rows is not None else load_prospects()
    if sector and sector.strip():
        term = sector.strip().lower()
        result = [r for r in result if term in (r.get("sector") or "").lower()]
    if priority is not None:
        result = [r for r in result if r.get("outreach_priority") == priority]
    if min_incentive_score is not None:
        result = [r for r in result if (r.get("incentive_score") or 0) >= min_incentive_score]
    if contact_type and contact_type.strip():
        term = contact_type.strip().lower()
        result = [r for r in result if term in (r.get("contact_type") or "").lower()]
    if q and q.strip():
        term = q.strip().lower()
        result = [
            r
            for r in result
            if term in (r.get("company") or "").lower()
            or term in (r.get("sector") or "").lower()
            or term in (r.get("why_attractive") or "").lower()
        ]
    if limit is not None and limit > 0:
        return result[:limit]
    return result


def get_prospect_meta() -> dict[str, Any]:
    """Alias for API routers."""
    return prospects_meta()


def _priority_component(priority: int) -> float:
    """Priority 1 (urgent) scores 100; priority 5 scores 20."""
    p = min(5, max(1, int(priority)))
    return ((6 - p) / 5.0) * 100.0


def _contact_type_match_score(row_contact_type: str, query: str | None) -> float:
    if not query or not query.strip():
        return 50.0
    a = query.strip().lower()
    b = (row_contact_type or "").lower()
    if a == b:
        return 100.0
    if a in b or b in a:
        return 85.0
    a_tokens = set(re.split(r"[\s,/]+", a))
    b_tokens = set(re.split(r"[\s,/]+", b))
    overlap = a_tokens & b_tokens
    if overlap:
        return 60.0 + min(30.0, len(overlap) * 10.0)
    return 25.0


def score_prospect(
    row: dict[str, Any],
    *,
    contact_type_query: str | None = None,
    weights: dict[str, float] | None = None,
) -> tuple[float, dict[str, float]]:
    w = {**_DEFAULT_WEIGHTS, **(weights or {})}
    total_w = sum(w.values()) or 1.0
    w = {k: v / total_w for k, v in w.items()}

    incentive = float(row.get("incentive_score") or 0)
    priority = _priority_component(int(row.get("outreach_priority") or 3))
    yale = 100.0 if row.get("has_yale_hook") else 0.0
    ctype = _contact_type_match_score(row.get("contact_type") or "", contact_type_query)

    breakdown = {
        "incentive": round(incentive, 1),
        "priority": round(priority, 1),
        "yale_hook": round(yale, 1),
        "contact_type_match": round(ctype, 1),
    }
    composite = (
        w["incentive"] * incentive
        + w["priority"] * priority
        + w["yale_hook"] * yale
        + w["contact_type_match"] * ctype
    )
    return round(composite, 2), breakdown


def build_verifiability_payload(
    row: dict[str, Any],
    *,
    composite_score: float,
    score_breakdown: dict[str, float],
) -> dict[str, Any]:
    return {
        "company": row.get("company"),
        "row_index": row.get("row_index"),
        "sector": row.get("sector"),
        "engagement_theme": row.get("engagement_theme"),
        "composite_score": composite_score,
        "score_breakdown": score_breakdown,
        "score_rationale": row.get("score_rationale"),
        "incentive_score": row.get("incentive_score"),
        "outreach_priority": row.get("outreach_priority"),
        "has_yale_hook": row.get("has_yale_hook"),
        "contact_type": row.get("contact_type"),
        "verification_source_url": row.get("verification_source_url"),
        "yucg_service_tags": row.get("yucg_service_tags") or [],
        "target_role_title": row.get("target_role_title"),
        "first_message_angle": row.get("first_message_angle"),
        "discovery_hint": row.get("discovery_hint"),
    }


def recommend_prospects(
    *,
    n: int = 10,
    sector: str | None = None,
    priority: int | None = None,
    min_incentive_score: float | None = None,
    contact_type: str | None = None,
    contact_type_match: str | None = None,
    q: str | None = None,
) -> list[dict[str, Any]]:
    rows = filter_prospects(
        load_prospects(),
        sector=sector,
        priority=priority,
        min_incentive_score=min_incentive_score,
        contact_type=contact_type,
        q=q,
    )
    type_for_score = contact_type_match or contact_type
    scored: list[tuple[float, dict[str, float], dict[str, Any]]] = []
    for row in rows:
        composite, breakdown = score_prospect(row, contact_type_query=type_for_score)
        scored.append((composite, breakdown, row))
    scored.sort(key=lambda x: (-x[0], x[2].get("row_index") or 0))

    out: list[dict[str, Any]] = []
    for composite, breakdown, row in scored[: max(1, n)]:
        breakdown_with_meta = {
            **breakdown,
            "total": composite,
            "rationale": row.get("score_rationale"),
        }
        verifiability = build_verifiability_payload(
            row,
            composite_score=composite,
            score_breakdown=breakdown_with_meta,
        )
        out.append(
            {
                "prospect": _prospect_api_row(row),
                "verifiability": verifiability,
                "composite_score": composite,
            }
        )
    return out


def _prospect_api_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "row_index": row.get("row_index"),
        "company": row.get("company"),
        "sector": row.get("sector"),
        "why_attractive": row.get("why_attractive"),
        "engagement_theme": row.get("engagement_theme"),
        "yale_hook": row.get("yale_hook"),
        "outreach_priority": row.get("outreach_priority"),
        "contact_type": row.get("contact_type"),
        "target_role_title": row.get("target_role_title"),
        "incentive_score": row.get("incentive_score"),
        "verification_source_url": row.get("verification_source_url"),
        "recommended_message_angle": row.get("first_message_angle"),
        "yucg_service_tags": row.get("yucg_service_tags") or [],
    }


def top_prospects_for_ai(
    *,
    limit: int = 30,
    sector: str | None = None,
    contact_type: str | None = None,
    min_incentive_score: float | None = None,
) -> list[dict[str, Any]]:
    """Top spreadsheet rows by incentive for Ollama prompt context."""
    rows = filter_prospects(
        sector=sector,
        contact_type=contact_type,
        min_incentive_score=min_incentive_score,
    )
    rows = sorted(
        rows,
        key=lambda r: (-(r.get("incentive_score") or 0), r.get("row_index") or 0),
    )
    return rows[: max(1, limit)]


def build_shortlist_csv(
    rows: list[dict[str, Any]] | None = None,
    *,
    row_indices: list[int] | None = None,
    sector: str | None = None,
    priority: int | None = None,
    min_incentive_score: float | None = None,
    contact_type: str | None = None,
    limit: int = 50,
) -> bytes:
    if rows is None:
        if row_indices:
            all_rows = {r["row_index"]: r for r in load_prospects()}
            rows = [all_rows[i] for i in row_indices if i in all_rows]
        else:
            rows = filter_prospects(
                sector=sector,
                priority=priority,
                min_incentive_score=min_incentive_score,
                contact_type=contact_type,
                limit=limit,
            )
            scored = []
            for row in rows:
                composite, _ = score_prospect(row, contact_type_query=contact_type)
                scored.append({**row, "composite_score": composite})
            rows = sorted(scored, key=lambda r: (-(r.get("composite_score") or 0), r.get("row_index") or 0))

    fieldnames = [
        "row_index",
        "company",
        "sector",
        "outreach_priority",
        "incentive_score",
        "contact_type",
        "target_role_title",
        "engagement_theme",
        "yale_hook",
        "verification_source_url",
        "first_message_angle",
        "discovery_hint",
        "yucg_service_tags",
        "composite_score",
    ]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for r in rows:
        tags = r.get("yucg_service_tags") or []
        writer.writerow(
            {
                **r,
                "yale_hook": r.get("yale_hook") or "",
                "yucg_service_tags": "; ".join(tags) if isinstance(tags, list) else str(tags),
                "composite_score": r.get("composite_score", ""),
            }
        )
    return buf.getvalue().encode("utf-8-sig")


def prospects_meta() -> dict[str, Any]:
    path = prospect_xlsx_path()
    rows = load_prospects()
    sectors = sorted({r["sector"] for r in rows if r.get("sector")})
    contact_types = sorted({r["contact_type"] for r in rows if r.get("contact_type")})
    return {
        "source_path": str(path),
        "source_exists": path.is_file(),
        "sheet": SHEET_NAME,
        "row_count": len(rows),
        "sectors": sectors,
        "contact_types": contact_types,
        "cached_mtime": _cache.get("mtime"),
    }
